#!/usr/bin/env python3
"""Convert the manually maintained Middle East strike Excel workbooks to JSON.

Supported inputs:
- USA–Iran strike workbook
- Houthi (Ansar Allah) strike workbook

Expected sheet in both workbooks: Események

The converter is backward compatible with the original Hungarian USA–Iran
workbook and also accepts the Houthi workbook column name:
    "Célország / térség"

Output:
- data/strike_history.json
- data/strike_history_summary.json
- data/strike_history_validation.json

Each output event contains:
    dataset_group = "USA_IRAN" or "HOUTHI"

This field is intended for dashboard toggles/filters.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# Column handling
# ---------------------------------------------------------------------------

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("Dátum",),
    "attacker": ("Támadó fél",),
    "country": ("Célország", "Célország / térség"),
    "location": ("Helyszín / célpont",),
    "latitude": ("Szélesség",),
    "longitude": ("Hosszúság",),
    "description": ("Esemény leírása",),
    "strike_type": ("Támadás típusa",),
    "coordinate_note": ("Koordináta / adat megjegyzés",),
    "confidence": ("Bizonyosság",),
    "source_url": ("Forrás URL",),
    # Optional fields:
    "source_name": ("Forrás",),
    "methodology_note": ("Módszertani megjegyzés",),
}


ATTACKER_MAP = {
    "usa": "USA",
    "egyesült államok": "USA",
    "united states": "USA",
    "us": "USA",
    "u.s.": "USA",

    "irán": "IRAN",
    "iran": "IRAN",
    "iranian": "IRAN",
    "irgc": "IRAN",

    "huszik": "HOUTHI",
    "huszi": "HOUTHI",
    "houthi": "HOUTHI",
    "houthis": "HOUTHI",
    "ansar allah": "HOUTHI",
    "ansarallah": "HOUTHI",
    "huszik (ansar allah)": "HOUTHI",
    "houthi (ansar allah)": "HOUTHI",
}


ATTACKER_LABELS = {
    "USA": "United States",
    "IRAN": "Iran",
    "HOUTHI": "Houthis / Ansar Allah",
}


CONFIDENCE_MAP = {
    "magas": "HIGH",
    "közepes": "MEDIUM",
    "alacsony": "LOW",
    "high": "HIGH",
    "medium": "MEDIUM",
    "low": "LOW",
}


DATASET_GROUPS = {
    "usa_iran": "USA_IRAN",
    "houthi": "HOUTHI",
}


# ---------------------------------------------------------------------------
# Normalization helpers
# ---------------------------------------------------------------------------

def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    text = clean_text(value)

    for fmt in (
        "%Y-%m-%d",
        "%Y.%m.%d",
        "%Y.%m.%d.",
        "%d.%m.%Y",
        "%d.%m.%Y.",
    ):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue

    raise ValueError(f"Nem értelmezhető dátum: {value!r}")


def normalize_attacker(value: Any) -> str:
    text = clean_text(value)
    return ATTACKER_MAP.get(text.casefold(), text.upper())


def normalize_confidence(value: Any) -> str:
    text = clean_text(value)
    return CONFIDENCE_MAP.get(text.casefold(), text.upper() or "UNKNOWN")


def normalize_coordinate(value: Any, field_name: str) -> float:
    if value is None or value == "":
        raise ValueError(f"Hiányzó koordináta: {field_name}")

    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"Hibás koordináta ({field_name}): {value!r}"
        ) from exc

    if not math.isfinite(number):
        raise ValueError(
            f"Nem véges koordináta ({field_name}): {value!r}"
        )

    if field_name == "latitude" and not -90 <= number <= 90:
        raise ValueError(f"Szélesség tartományon kívül: {number}")

    if field_name == "longitude" and not -180 <= number <= 180:
        raise ValueError(f"Hosszúság tartományon kívül: {number}")

    return round(number, 6)


def resolve_columns(headers: list[str]) -> dict[str, int]:
    """Resolve canonical field names against accepted Hungarian aliases."""

    resolved: dict[str, int] = {}

    for canonical_name, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias in headers:
                resolved[canonical_name] = headers.index(alias)
                break

    required = {
        "date",
        "attacker",
        "country",
        "location",
        "latitude",
        "longitude",
        "description",
        "strike_type",
        "coordinate_note",
        "confidence",
        "source_url",
    }

    missing = sorted(required - set(resolved))

    if missing:
        readable = []
        for field in missing:
            aliases = " / ".join(HEADER_ALIASES[field])
            readable.append(aliases)

        raise RuntimeError(
            "Hiányzó kötelező oszlopok: " + ", ".join(readable)
        )

    return resolved


def get_optional_value(
    row: tuple[Any, ...],
    column: dict[str, int],
    field_name: str,
) -> str:
    index = column.get(field_name)
    if index is None:
        return ""
    return clean_text(row[index])


# ---------------------------------------------------------------------------
# Event IDs
# ---------------------------------------------------------------------------

def attacker_code(attacker: str) -> str:
    if attacker == "USA":
        return "US"
    if attacker == "IRAN":
        return "IR"
    if attacker == "HOUTHI":
        return "HU"
    return "OT"


def make_event_id(
    event_date: str,
    attacker: str,
    country: str,
    location: str,
    source_file: str,
    row_number: int,
) -> str:
    # Including source_file avoids collisions when two workbooks contain
    # similarly dated events.
    raw = (
        f"{event_date}|{attacker}|{country}|{location}|"
        f"{source_file}|{row_number}"
    )
    digest = hashlib.sha1(
        raw.encode("utf-8")
    ).hexdigest()[:8].upper()

    return (
        f"ME-{event_date.replace('-', '')}-"
        f"{attacker_code(attacker)}-{digest}"
    )


# ---------------------------------------------------------------------------
# Workbook parser
# ---------------------------------------------------------------------------

def read_events(
    input_path: Path,
    dataset_group: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:

    workbook = load_workbook(
        input_path,
        data_only=True,
        read_only=True,
    )

    if "Események" not in workbook.sheetnames:
        raise RuntimeError(
            f"{input_path.name}: az Excel nem tartalmaz "
            "'Események' nevű munkalapot."
        )

    sheet = workbook["Események"]

    headers = [
        clean_text(cell.value)
        for cell in next(
            sheet.iter_rows(min_row=1, max_row=1)
        )
    ]

    column = resolve_columns(headers)

    events: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []

    for row_number, row in enumerate(
        sheet.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        if not any(
            value not in (None, "")
            for value in row
        ):
            continue

        try:
            event_date = normalize_date(
                row[column["date"]]
            )

            attacker = normalize_attacker(
                row[column["attacker"]]
            )

            country = clean_text(
                row[column["country"]]
            )

            location = clean_text(
                row[column["location"]]
            )

            latitude = normalize_coordinate(
                row[column["latitude"]],
                "latitude",
            )

            longitude = normalize_coordinate(
                row[column["longitude"]],
                "longitude",
            )

            description = clean_text(
                row[column["description"]]
            )

            strike_type = clean_text(
                row[column["strike_type"]]
            )

            coordinate_note = clean_text(
                row[column["coordinate_note"]]
            )

            confidence_hu = clean_text(
                row[column["confidence"]]
            )

            source_url = clean_text(
                row[column["source_url"]]
            )

            source_name = get_optional_value(
                row,
                column,
                "source_name",
            )

            methodology_note = get_optional_value(
                row,
                column,
                "methodology_note",
            )

            required_text = {
                "Támadó fél": attacker,
                "Célország / térség": country,
                "Helyszín / célpont": location,
                "Esemény leírása": description,
                "Támadás típusa": strike_type,
                "Forrás URL": source_url,
            }

            empty_fields = [
                name
                for name, value in required_text.items()
                if not value
            ]

            if empty_fields:
                raise ValueError(
                    "Hiányzó kötelező mező(k): "
                    + ", ".join(empty_fields)
                )

            known_attackers = {
                "USA",
                "IRAN",
                "HOUTHI",
            }

            if attacker not in known_attackers:
                warnings.append({
                    "source_file": input_path.name,
                    "dataset_group": dataset_group,
                    "row": row_number,
                    "warning":
                        f"Ismeretlen támadó fél: {attacker}",
                })

            if (
                source_url
                and not source_url.startswith(
                    ("http://", "https://")
                )
            ):
                warnings.append({
                    "source_file": input_path.name,
                    "dataset_group": dataset_group,
                    "row": row_number,
                    "warning":
                        "A forrás URL nem http/https címmel kezdődik.",
                })

            event = {
                "event_id": make_event_id(
                    event_date,
                    attacker,
                    country,
                    location,
                    input_path.name,
                    row_number,
                ),

                # Dataset-level filter for the dashboard.
                "dataset_group": dataset_group,

                "date": event_date,

                "attacker": attacker,
                "attacker_label":
                    ATTACKER_LABELS.get(
                        attacker,
                        attacker,
                    ),

                "target_country": country,
                "target_location": location,

                # Aliases used by some existing map code.
                "country": country,
                "location": location,

                "latitude": latitude,
                "longitude": longitude,

                # Leaflet compatibility.
                "lat": latitude,
                "lon": longitude,

                "description": description,
                "strike_type": strike_type,
                "coordinate_note": coordinate_note,

                "confidence":
                    normalize_confidence(
                        confidence_hu
                    ),

                "confidence_label_hu":
                    confidence_hu,

                # iranstrike-layer.js can already consume this field name.
                "attacker_confidence":
                    normalize_confidence(
                        confidence_hu
                    ).lower(),

                "source_name":
                    source_name
                    or input_path.stem,

                "source_url": source_url,

                "methodology_note":
                    methodology_note,

                "source_file":
                    input_path.name,

                "source_row":
                    row_number,

                # Explicit manual-event metadata.
                "data_origin":
                    "manual_excel",

                "map_visualizable":
                    True,
            }

            events.append(event)

        except Exception as exc:
            warnings.append({
                "source_file": input_path.name,
                "dataset_group": dataset_group,
                "row": row_number,
                "error": str(exc),
            })

    events.sort(
        key=lambda item: (
            item["date"],
            item["attacker"],
            item["event_id"],
        )
    )

    return events, warnings


# ---------------------------------------------------------------------------
# Analytics
# ---------------------------------------------------------------------------

def build_summary(
    events: list[dict[str, Any]],
) -> dict[str, Any]:

    by_attacker = Counter(
        event["attacker"]
        for event in events
    )

    by_country = Counter(
        event["target_country"]
        for event in events
    )

    by_date = Counter(
        event["date"]
        for event in events
    )

    by_dataset = Counter(
        event["dataset_group"]
        for event in events
    )

    by_type = Counter(
        event["strike_type"]
        for event in events
    )

    dates = [
        event["date"]
        for event in events
    ]

    return {
        "event_count":
            len(events),

        "date_start":
            min(dates) if dates else None,

        "date_end":
            max(dates) if dates else None,

        "datasets":
            dict(
                sorted(
                    by_dataset.items()
                )
            ),

        "attackers":
            dict(
                sorted(
                    by_attacker.items()
                )
            ),

        "target_countries":
            dict(
                sorted(
                    by_country.items(),
                    key=lambda x: (
                        -x[1],
                        x[0],
                    ),
                )
            ),

        "strike_types":
            dict(
                sorted(
                    by_type.items(),
                    key=lambda x: (
                        -x[1],
                        x[0],
                    ),
                )
            ),

        "daily_counts":
            dict(
                sorted(
                    by_date.items()
                )
            ),
    }


def write_json(
    path: Path,
    payload: dict[str, Any],
) -> None:

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "USA–Irán + huszi csapásadatok "
            "Excel → közös JSON feldolgozó"
        )
    )

    parser.add_argument(
        "--input",
        default=(
            "data/manual/"
            "usa_iran_tamadasok.xlsx"
        ),
        help="USA–Irán Excel",
    )

    parser.add_argument(
        "--houthi-input",
        default=(
            "data/manual/"
            "huszi_tamadasok.xlsx"
        ),
        help="Huszi Excel",
    )

    parser.add_argument(
        "--output",
        default=(
            "data/"
            "strike_history.json"
        ),
    )

    parser.add_argument(
        "--summary",
        default=(
            "data/"
            "strike_history_summary.json"
        ),
    )

    parser.add_argument(
        "--validation",
        default=(
            "data/"
            "strike_history_validation.json"
        ),
    )

    parser.add_argument(
        "--skip-houthi",
        action="store_true",
        help=(
            "Csak a USA–Irán Excel feldolgozása. "
            "Visszafelé kompatibilis futtatás."
        ),
    )

    args = parser.parse_args()

    input_specs: list[tuple[Path, str]] = [
        (
            Path(args.input),
            DATASET_GROUPS["usa_iran"],
        )
    ]

    if not args.skip_houthi:
        input_specs.append(
            (
                Path(args.houthi_input),
                DATASET_GROUPS["houthi"],
            )
        )

    missing_files = [
        str(path)
        for path, _dataset_group
        in input_specs
        if not path.exists()
    ]

    if missing_files:
        print(
            "HIBA: Nem található bemeneti Excel:",
            file=sys.stderr,
        )

        for filename in missing_files:
            print(
                f"  - {filename}",
                file=sys.stderr,
            )

        return 1

    all_events: list[dict[str, Any]] = []
    all_warnings: list[dict[str, Any]] = []
    source_files: list[str] = []

    for input_path, dataset_group in input_specs:
        events, warnings = read_events(
            input_path,
            dataset_group,
        )

        all_events.extend(events)
        all_warnings.extend(warnings)
        source_files.append(input_path.name)

        print(
            f"{dataset_group}: "
            f"{len(events)} esemény "
            f"({input_path.name})"
        )

    all_events.sort(
        key=lambda item: (
            item["date"],
            item["dataset_group"],
            item["attacker"],
            item["event_id"],
        )
    )

    errors = [
        item
        for item in all_warnings
        if "error" in item
    ]

    generated_at = (
        datetime.now(timezone.utc)
        .replace(microsecond=0)
        .isoformat()
        .replace("+00:00", "Z")
    )

    summary = build_summary(
        all_events
    )

    dataset = {
        "generated_at":
            generated_at,

        "dataset":
            "middle_east_strike_history",

        "dataset_version":
            2,

        "source_files":
            source_files,

        "available_dataset_groups": [
            "USA_IRAN",
            "HOUTHI",
        ],

        "summary":
            summary,

        "events":
            all_events,
    }

    write_json(
        Path(args.output),
        dataset,
    )

    write_json(
        Path(args.summary),
        {
            "generated_at":
                generated_at,

            **summary,
        },
    )

    write_json(
        Path(args.validation),
        {
            "generated_at":
                generated_at,

            "source_files":
                source_files,

            "valid_event_count":
                len(all_events),

            "error_count":
                len(errors),

            "warning_count":
                len(all_warnings)
                - len(errors),

            "items":
                all_warnings,
        },
    )

    print(
        f"Összes feldolgozott esemény: "
        f"{len(all_events)}"
    )

    print(
        f"Hibák: {len(errors)} | "
        f"Figyelmeztetések: "
        f"{len(all_warnings) - len(errors)}"
    )

    print(
        "Dataset csoportok: "
        + ", ".join(
            f"{key}={value}"
            for key, value
            in summary["datasets"].items()
        )
    )

    if errors:
        print(
            "HIBA: Egy vagy több Excel-sor "
            "nem volt feldolgozható.",
            file=sys.stderr,
        )
        return 2

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
